"""Bounded file parsing; intentionally no student or regulatory validation."""

import asyncio
import csv
import io
import re
import sys
import zipfile
from collections.abc import AsyncGenerator, Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import PurePath
from time import monotonic
from typing import Any, Literal

from fastapi import Request
from openpyxl import load_workbook
from pydantic import ValidationError
from starlette.datastructures import FormData, UploadFile
from starlette.formparsers import MultiPartException, MultiPartParser
from starlette.requests import ClientDisconnect

from app.core.config import Settings
from app.core.http import APIError
from app.modules.submissions.schemas import UploadContext

HEADERS = (
    "college_student_id",
    "student_name",
    "apaar_id",
    "academic_year",
    "programme_code",
    "admission_category",
    "enrollment_status",
    "enrollment_date",
)
FORM_FIELDS = {
    "institution_id",
    "programme_id",
    "academic_year",
    "schema_version",
    "file",
}
ZIP_MAGIC = b"PK\x03\x04"
UNSUPPORTED_EXTENSIONS = {".xlsm", ".xlsb", ".xls", ".xltx", ".xltm", ".xlam"}
XLSX_MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
XLSX_MAX_MEMBERS = 500
XLSX_BLANK_ROW_ALLOWANCE = 100
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Actual request bytes are bounded first. Avoid an unrelated global 128 KiB cell cap.
csv.field_size_limit(sys.maxsize)


def check_budget(deadline: float) -> None:
    if monotonic() >= deadline:
        raise APIError(408, "UPLOAD_TIMEOUT", "Upload processing budget expired")


@dataclass(frozen=True, repr=False)
class ParsedRow:
    number: int
    original: dict[str, str]
    normalized: dict[str, str]
    enrollment_date: date | None


def collect_rows(
    records: Iterable[Sequence[str]],
    settings: Settings,
    deadline: float,
    *,
    label: str,
    error_code: str,
) -> list[ParsedRow]:
    """Shared header check, limits and normalization for every upload format."""
    iterator = iter(records)
    if tuple(next(iterator, ())) != HEADERS:
        raise APIError(422, "SCHEMA_MISMATCH", f"{label} headers must match v1")
    rows: list[ParsedRow] = []
    for number, values in enumerate(iterator, start=2):
        check_budget(deadline)
        if len(values) != len(HEADERS):
            raise APIError(
                400, error_code, f"{label} record column count differs from v1"
            )
        if len(rows) >= settings.upload_max_rows:
            raise APIError(
                422, "ROW_LIMIT_EXCEEDED", f"{label} exceeds configured record limit"
            )
        original = dict(zip(HEADERS, values, strict=True))
        normalized = {key: value.strip() for key, value in original.items()}
        for key in ("programme_code", "admission_category", "enrollment_status"):
            normalized[key] = normalized[key].upper()
        parsed_date = None
        if re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", normalized["enrollment_date"]):
            try:
                parsed_date = date.fromisoformat(normalized["enrollment_date"])
            except ValueError:
                pass
        rows.append(ParsedRow(number, original, normalized, parsed_date))
    if not rows:
        raise APIError(400, error_code, f"{label} has no data records")
    check_budget(deadline)
    return rows


def parse_csv(data: bytes, settings: Settings, deadline: float) -> list[ParsedRow]:
    if not data:
        raise APIError(400, "INVALID_CSV", "CSV has no data records")
    if len(data) > settings.upload_max_bytes:
        raise APIError(413, "FILE_TOO_LARGE", "File exceeds configured limit")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise APIError(400, "INVALID_CSV", "CSV must use UTF-8") from None
    if "\x00" in text:
        raise APIError(400, "INVALID_CSV", "Invalid CSV content")
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)
    try:
        return collect_rows(
            reader, settings, deadline, label="CSV", error_code="INVALID_CSV"
        )
    except csv.Error:
        raise APIError(400, "INVALID_CSV", "Malformed CSV") from None


def detect_format(data: bytes, filename: str) -> Literal["csv", "xlsx"]:
    """Extension and content signature must agree; neither is trusted alone."""
    extension = PurePath(filename).suffix.lower()
    is_zip = data.startswith(ZIP_MAGIC)
    if extension in UNSUPPORTED_EXTENSIONS:
        raise APIError(400, "UNSUPPORTED_FILE_TYPE", "Unsupported file type")
    if extension == ".xlsx":
        if not is_zip:
            raise APIError(400, "INVALID_XLSX", "File is not a valid xlsx workbook")
        return "xlsx"
    if is_zip:
        raise APIError(400, "UNSUPPORTED_FILE_TYPE", "File type does not match name")
    return "csv"


def check_xlsx_package(data: bytes) -> None:
    """Reject non-xlsx, macro-enabled, encrypted or expanding packages up front."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            infos = package.infolist()
            names = {info.filename for info in infos}
            if len(infos) > XLSX_MAX_MEMBERS or (
                sum(info.file_size for info in infos) > XLSX_MAX_UNCOMPRESSED_BYTES
            ):
                raise APIError(413, "FILE_TOO_LARGE", "Workbook expands beyond limit")
            if any(info.flag_bits & 0x1 for info in infos):
                raise APIError(400, "INVALID_XLSX", "Encrypted workbooks unsupported")
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                raise APIError(400, "INVALID_XLSX", "File is not a valid xlsx workbook")
            content_types = package.read("[Content_Types].xml")
            if (
                b"macroEnabled" in content_types
                or b"vnd.ms-office.vbaProject" in content_types
                or any(name.lower().startswith("xl/vbaproject") for name in names)
            ):
                raise APIError(400, "INVALID_XLSX", "Macro-enabled workbooks rejected")
            if b"spreadsheetml.sheet.main+xml" not in content_types:
                raise APIError(400, "INVALID_XLSX", "File is not a valid xlsx workbook")
            for info in infos:
                if info.filename.endswith((".xml", ".rels")):
                    with package.open(info) as member:
                        if b"<!ENTITY" in member.read():
                            raise APIError(
                                400, "INVALID_XLSX", "Unsafe workbook content"
                            )
    except (zipfile.BadZipFile, zipfile.LargeZipFile, OSError, KeyError, EOFError):
        raise APIError(
            400, "INVALID_XLSX", "File is not a valid xlsx workbook"
        ) from None


def cell_text(value: object) -> str:
    """Render a cell as the string the CSV path would have carried."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def xlsx_records(
    sheet: Any, settings: Settings, deadline: float
) -> Iterator[list[str]]:
    # Declared dimensions are untrusted metadata; the scan below is bounded too.
    scan_limit = settings.upload_max_rows + 1 + XLSX_BLANK_ROW_ALLOWANCE
    if (sheet.max_row or 0) > scan_limit:
        raise APIError(
            422, "ROW_LIMIT_EXCEEDED", "Workbook exceeds configured record limit"
        )
    width = len(HEADERS)
    blanks = 0
    for scanned, cells in enumerate(
        sheet.iter_rows(max_col=width + 1, values_only=True), start=1
    ):
        check_budget(deadline)
        if scanned > scan_limit:
            raise APIError(
                422, "ROW_LIMIT_EXCEEDED", "Workbook exceeds configured record limit"
            )
        values = [cell_text(cell) for cell in cells]
        while values and not values[-1].strip():
            values.pop()
        if not values and scanned > 1:
            blanks += 1  # Trailing blank rows are formatting leftovers; drop later.
            continue
        for _ in range(blanks):  # Interior blanks keep their source row numbers.
            yield [""] * width
        blanks = 0
        yield values + [""] * (width - len(values))


def parse_xlsx(data: bytes, settings: Settings, deadline: float) -> list[ParsedRow]:
    if len(data) > settings.upload_max_bytes:
        raise APIError(413, "FILE_TOO_LARGE", "File exceeds configured limit")
    check_xlsx_package(data)
    workbook = None
    try:
        # data_only returns cached values; formulas are never evaluated.
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
        if not workbook.worksheets:
            raise APIError(400, "INVALID_XLSX", "Workbook has no worksheet")
        return collect_rows(
            xlsx_records(workbook.worksheets[0], settings, deadline),
            settings,
            deadline,
            label="Workbook",
            error_code="INVALID_XLSX",
        )
    except APIError:
        raise
    except Exception:
        raise APIError(400, "INVALID_XLSX", "Malformed workbook") from None
    finally:
        if workbook is not None:
            workbook.close()


class CompleteMultipartParser(MultiPartParser):
    """Detect truncation while retaining Starlette's exception cleanup."""

    complete = False

    def on_end(self) -> None:
        self.complete = True


async def read_multipart(
    request: Request, settings: Settings
) -> tuple[UploadContext, bytes, str, float]:
    limit = settings.upload_max_bytes + 64 * 1024
    length = request.headers.get("content-length")
    if length is not None:
        try:
            size = int(length)
            if size < 0:
                raise ValueError
            if size > limit:
                raise APIError(
                    413, "FILE_TOO_LARGE", "Request exceeds configured limit"
                )
        except ValueError:
            raise APIError(400, "INVALID_CSV", "Invalid request length") from None
    if (
        request.headers.get("content-type", "").split(";")[0].strip().lower()
        != "multipart/form-data"
    ):
        raise APIError(400, "INVALID_CSV", "Expected multipart upload")
    body = bytearray()
    try:
        async with asyncio.timeout(settings.upload_receive_timeout_seconds):
            async for chunk in request.stream():
                if len(body) + len(chunk) > limit:
                    raise APIError(
                        413, "FILE_TOO_LARGE", "Request exceeds configured limit"
                    )
                body.extend(chunk)
    except TimeoutError:
        raise APIError(408, "UPLOAD_TIMEOUT", "Upload receive budget expired") from None
    except ClientDisconnect:
        raise APIError(400, "INVALID_CSV", "Upload disconnected") from None
    deadline = monotonic() + settings.upload_processing_timeout_seconds

    async def stream() -> AsyncGenerator[bytes]:
        yield bytes(body)
        # Raising inside parse makes Starlette close even unfinished file parts.
        if not parser.complete:
            raise MultiPartException("Incomplete multipart body")

    parser = CompleteMultipartParser(
        request.headers, stream(), max_files=1, max_fields=4, max_part_size=4096
    )
    # Within the already bounded body no file rolls out to the OS temporary folder.
    parser.spool_max_size = limit + 1
    form: FormData | None = None
    try:
        form = await parser.parse()
        check_budget(deadline)
        items = form.multi_items()
        keys = [key for key, _ in items]
        if set(keys) - FORM_FIELDS or len(keys) != len(set(keys)):
            raise APIError(
                422, "INVALID_REQUEST", "Unexpected or duplicate upload fields"
            )
        values = {key: value for key, value in items if key != "file"}
        try:
            context = UploadContext.model_validate(values)
        except ValidationError:
            raise APIError(422, "INVALID_REQUEST", "Invalid upload context") from None
        file = form.get("file")
        if not isinstance(file, UploadFile) or not file.filename:
            raise APIError(422, "INVALID_REQUEST", "One CSV or XLSX file is required")
        if len(file.filename) > 255 or any(ord(char) < 32 for char in file.filename):
            raise APIError(422, "INVALID_REQUEST", "Invalid file metadata")
        if (file.content_type or "").split(";")[0].strip().lower() not in {
            "text/csv",
            "application/csv",
            "text/plain",
            "application/octet-stream",
            "application/vnd.ms-excel",
            XLSX_CONTENT_TYPE,
        }:
            raise APIError(400, "INVALID_CSV", "Unsupported file type")
        data = await file.read(settings.upload_max_bytes + 1)
        if len(data) > settings.upload_max_bytes:
            raise APIError(413, "FILE_TOO_LARGE", "File exceeds configured limit")
        check_budget(deadline)
        return context, data, file.filename, deadline
    except MultiPartException as exc:
        if exc.message.startswith(
            ("Too many fields", "Too many files", "Part exceeded")
        ):
            raise APIError(422, "INVALID_REQUEST", "Invalid multipart fields") from None
        raise APIError(400, "INVALID_CSV", "Malformed multipart upload") from None
    finally:
        if form is not None:
            await form.close()
