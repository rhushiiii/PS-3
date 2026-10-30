"""CSV, XLSX and PDF renderers for a ReportDocument.

CSV layout (UTF-8, no BOM, CRLF): leading comment-style metadata rows whose first
cell starts with "# " (a key in column 1, the value in column 2), then for each of
the two tables a "# <table title>" row, a header row and the data rows, with a blank
row between tables. Readers should skip rows beginning with "#". A missing value is
an empty cell. Every text cell passes through the shared formula-injection guard.
"""

import csv
import io
import re
from collections.abc import Callable

from fpdf import FPDF
from fpdf.enums import Align, XPos, YPos
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.reports.schemas import (
    Cell,
    ReportDocument,
    ReportFormat,
    ReportTable,
)
from app.modules.validation.services import safe_csv_cell

MEDIA_TYPES: dict[str, str] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
NO_DATA = "No data"
UNICODE_TO_LATIN1 = str.maketrans(
    {
        "—": "-",
        "–": "-",
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "…": "...",
    }
)


def text_of(value: Cell) -> str:
    """Plain text for a cell; None is handled by each renderer."""
    if isinstance(value, bool):
        return "true" if value else "false"
    return "" if value is None else str(value)


# ------------------------------------------------------------------------ CSV


def to_csv(report: ReportDocument) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)

    def row(cells: list[Cell]) -> None:
        # Numbers are written as-is (a leading "-" is not a formula); text is guarded.
        writer.writerow(
            [safe_csv_cell(c) if isinstance(c, str) else text_of(c) for c in cells]
        )

    for key, value in report.metadata():
        row([f"# {key}", value])
    for section in (report.totals, report.data):
        row([])
        row([f"# {section.title}"])
        row(list(section.columns))
        for cells in section.rows:
            row(cells)
    return buffer.getvalue().encode("utf-8")


# ------------------------------------------------------------------------ XLSX


def to_xlsx(report: ReportDocument) -> bytes:
    workbook = Workbook()
    meta = workbook.active
    assert meta is not None
    meta.title = "Metadata"
    data = workbook.create_sheet("Data")
    bold = Font(bold=True)

    def write_row(
        sheet: Worksheet, number: int, cells: list[Cell], *, head: bool
    ) -> None:
        for column, value in enumerate(cells, start=1):
            cell = sheet.cell(row=number, column=column)
            if isinstance(value, str):
                # Force a text cell so "=..." can never become a formula.
                cell.value = safe_csv_cell(value)
                cell.data_type = "s"
            else:
                cell.value = value
            if head:
                cell.font = bold

    for number, (key, value) in enumerate(report.metadata(), start=1):
        write_row(meta, number, [key, value], head=False)
        meta.cell(row=number, column=1).font = bold
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 100

    number = 1
    for section in (report.totals, report.data):
        write_row(data, number, [section.title], head=True)
        write_row(data, number + 1, list(section.columns), head=True)
        number += 2
        for cells in section.rows:
            write_row(data, number, cells, head=False)
            number += 1
        number += 1
    for index in range(
        1, max(len(report.data.columns), len(report.totals.columns)) + 1
    ):
        data.column_dimensions[get_column_letter(index)].width = 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# ------------------------------------------------------------------------- PDF


def latin1(text: str) -> str:
    """Core PDF fonts are Latin-1 only; map common punctuation, replace the rest."""
    return (
        text.translate(UNICODE_TO_LATIN1).encode("latin-1", "replace").decode("latin-1")
    )


class ReportPDF(FPDF):
    def __init__(self, provenance: str) -> None:
        super().__init__(orientation="L", unit="mm", format="A4")
        self.provenance = latin1(provenance)
        self.set_auto_page_break(auto=True, margin=20)
        self.alias_nb_pages()

    def footer(self) -> None:
        self.set_y(-17)
        self.set_font("Helvetica", "I", 7)
        self.multi_cell(
            0,
            3.5,
            self.provenance,
            align=Align.C,
            new_x=XPos.LMARGIN,
            new_y=YPos.NEXT,
        )
        self.cell(0, 4, f"Page {self.page_no()} of {{nb}}", align=Align.C)


def column_weights(section: ReportTable) -> tuple[int, ...]:
    """Relative widths from the longest word per column, so no word is forced narrow."""
    weights = []
    for index, column in enumerate(section.columns):
        words = column.replace("_", " ").split()
        words += [w for r in section.rows for w in text_of(r[index]).split()]
        weights.append(min(max((len(w) for w in words), default=4) + 2, 30))
    return tuple(weights)


def pdf_table(pdf: ReportPDF, section: ReportTable) -> None:
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(0, 6, latin1(section.title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 7)
    if not section.rows:
        pdf.cell(0, 5, NO_DATA, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(2)
        return
    with pdf.table(
        col_widths=column_weights(section),
        text_align="LEFT",
        line_height=3.8,
        headings_style=FontFace(emphasis="BOLD", fill_color=(225, 225, 225)),
    ) as table:
        header = table.row()
        for column in section.columns:
            header.cell(latin1(column.replace("_", " ")))
        for cells in section.rows:
            row = table.row()
            for value in cells:
                row.cell(latin1(NO_DATA if value is None else text_of(value)))
    pdf.ln(3)


def to_pdf(report: ReportDocument) -> bytes:
    pdf = ReportPDF(report.provenance)
    pdf.set_title(latin1(report.title))
    pdf.set_creation_date(report.generated_at)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 7, latin1(report.title), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 8)
    for key, value in report.metadata():
        pdf.multi_cell(
            0, 4, latin1(f"{key}: {value}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT
        )
    pdf.ln(3)
    pdf_table(pdf, report.totals)
    pdf_table(pdf, report.data)
    return bytes(pdf.output())


RENDERERS: dict[str, Callable[[ReportDocument], bytes]] = {
    "csv": to_csv,
    "xlsx": to_xlsx,
    "pdf": to_pdf,
}


def render(report: ReportDocument, file_format: ReportFormat) -> bytes:
    return RENDERERS[file_format](report)


SAFE_NAME = re.compile(r"[^a-z0-9-]")


def filename(report: ReportDocument, file_format: ReportFormat) -> str:
    """Generated only from the fixed kind, the clock and the format."""
    stamp = report.generated_at.strftime("%Y%m%dT%H%M%SZ")
    return f"{SAFE_NAME.sub('', report.kind)}-{stamp}.{file_format}"
