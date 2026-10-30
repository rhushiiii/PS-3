"""Phase 10: the scripted 12-step demonstration, measured against known truth.

Everything runs through the public HTTP API with seeded DEMO accounts on synthetic
files from samples/demo. The ground truth is samples/demo/expected.json.
"""

import io
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import httpx2
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import demo_scenario
from app.core.config import Settings
from app.demo_scenario import (
    SAMPLES,
    Scenario,
    build_xlsx,
    load_manifest,
    new_signing_key,
    raw_apaar_values,
)
from app.modules.audit.models import AuditEvent
from app.modules.intake.models import IntakeAssessment
from app.modules.submissions.parsing import HEADERS
from app.modules.validation.models import ValidationFinding

MANIFEST = load_manifest()
INTAKE_EXCEEDED_TEXT = (
    "candidate ACTIVE 4 > approved 3 (rule DEMO-v1, REGULAR) = excess 1"
)
SHORTFALL_TEXT = "candidate records 4 < declared 6 (DEMO declaration) = shortfall 2"


# ------------------------------------------------------------- no database needed


def test_sample_files_match_manifest() -> None:
    for name in MANIFEST["source_csv"]:
        text = (SAMPLES / name).read_text(encoding="utf-8")
        header, *rows = text.splitlines()
        assert header.split(",") == list(HEADERS), name
        assert not text.startswith("﻿")
        entry = MANIFEST["files"].get(name)
        if entry is not None:
            assert len(rows) == entry["rows"], name
    for name, entry in MANIFEST["files"].items():
        if "generated_from" in entry:
            assert not (SAMPLES / name).exists(), (
                "xlsx must be generated, not committed"
            )
            assert (
                len(
                    demo_scenario.read_csv_rows(
                        (SAMPLES / entry["generated_from"]).read_bytes()
                    )
                )
                == entry["rows"]
            )


def test_apaar_like_values_are_obviously_synthetic() -> None:
    for value in raw_apaar_values(MANIFEST):
        well_formed = re.fullmatch(r"[0-9]{12}", value) is not None
        assert (well_formed and value.startswith("9000")) or not well_formed, value
    people = " ".join(
        row["student_name"]
        for name in MANIFEST["source_csv"]
        for row in demo_scenario.read_csv_rows((SAMPLES / name).read_bytes())
    )
    assert "Demo Learner" in people and "@" not in people
    assert MANIFEST["label"].startswith("SIMULATED / DEMO DATA")


def test_generated_workbook_matches_source_rows() -> None:
    source = (SAMPLES / "08_willow_valid.csv").read_bytes()
    sheet = load_workbook(io.BytesIO(build_xlsx(source))).worksheets[0]
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    expected = demo_scenario.read_csv_rows(source)
    assert values[0] == list(HEADERS)
    assert [dict(zip(HEADERS, row, strict=True)) for row in values[1:]] == expected


def test_runner_refuses_non_local_and_non_loopback(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.setenv("DATABASE_URL", "postgresql+psycopg://x@127.0.0.1:1/none")
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.setenv("DEMO_AUTH", "false")
    assert demo_scenario.main([]) == 2
    assert "local DEMO_AUTH mode is required" in capsys.readouterr().err
    assert demo_scenario.main(["--base-url", "https://example.org"]) == 2
    assert "loopback" in capsys.readouterr().err


def test_ratio_helper_is_total() -> None:
    assert demo_scenario.ratio(0, 0) == 1.0
    assert demo_scenario.ratio(9, 9) == 1.0
    assert demo_scenario.ratio(1, 4) == 0.25


# ------------------------------------------------------------- full scenario


@pytest.fixture
def signed_client(demo_settings: Settings, client: TestClient) -> TestClient:
    # The Ed25519 key lives only in this test's settings object (ephemeral).
    demo_settings.credential_signing_key = SecretStr(new_signing_key())
    return client


def expected_counts(name: str) -> Counter[str]:
    return Counter(f["rule_code"] for f in MANIFEST["files"][name]["findings"])


def test_demo_scenario_end_to_end(
    signed_client: TestClient,
    seeded_session: Session,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    scenario = Scenario(signed_client)
    scenario.run()
    summary = scenario.summary()
    out = tmp_path / "demo_e2e_summary.json"
    out.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    with capsys.disabled():
        print("\n" + demo_scenario.render_table(summary))
        print(f"machine-readable summary: {out}")
        print(
            "ACCURACY_REPORT "
            + json.dumps(
                {
                    k: summary[k]
                    for k in (
                        "workflow",
                        "validation_accuracy",
                        "intake_correctness",
                        "integrity_signals",
                    )
                }
                | {
                    "reconciliation": {
                        "equal": summary["reconciliation"]["equal"],
                        "total": summary["reconciliation"]["total"],
                    },
                    "authorization": {
                        "passed": summary["authorization"]["passed"],
                        "total": summary["authorization"]["total"],
                    },
                    "raw_apaar_leaks": len(summary["raw_apaar_in_responses"]["leaks"]),
                }
            )
        )

    # Workflow completion: 12/12, every expectation inside every step.
    failures = [
        (s["number"], s["error"], s["failed_checks"])
        for s in summary["steps"]
        if s["status"] != "PASS"
    ]
    assert failures == []
    assert summary["workflow"] == {"steps_passed": 12, "steps_total": 12}
    assert summary["fresh_database"] is True
    assert json.loads(out.read_text())["passed"] is True

    # Validation precision/recall against the manifest (deterministic rules).
    accuracy = summary["validation_accuracy"]
    expected_total = sum(len(e["findings"]) for e in MANIFEST["files"].values())
    assert expected_total == 9
    assert (accuracy["true_positives"], accuracy["false_positives"]) == (9, 0)
    assert accuracy["false_negatives"] == 0
    assert accuracy["precision"] == accuracy["recall"] == 1.0
    assert accuracy["carryover_info_findings"] == 0
    assert accuracy["evidence_checks"] == 3 and accuracy["evidence_mismatches"] == 0
    for name, stats in accuracy["per_file"].items():
        assert stats["produced"] == stats["expected"] == stats["tp"], name

    # Independent oracle: the stored findings, straight from the database.
    stored: dict[str, Counter[str]] = {}
    for name, batch_id in scenario.batches.items():
        codes = seeded_session.scalars(
            select(ValidationFinding.rule_code).where(
                ValidationFinding.batch_id == batch_id
            )
        )
        stored[name] = Counter(codes)
        assert stored[name] == expected_counts(name), name
    assert sum(sum(c.values()) for c in stored.values()) == expected_total

    # Intake: exact excess and calculation text, also read back from the database.
    assert (
        summary["intake_correctness"]["exact"]
        == summary["intake_correctness"]["checks"]
    )
    assert summary["intake_correctness"]["checks"] > 40
    over = seeded_session.scalar(
        select(IntakeAssessment).where(
            IntakeAssessment.batch_id
            == scenario.batches["05_cedar_intake_excess_and_peer_match.csv"]
        )
    )
    assert over is not None
    regular = next(c for c in over.result["categories"] if c["category"] == "REGULAR")
    assert (regular["approved_seats"], regular["candidate_active_count"]) == (3, 4)
    assert (regular["exceeded"], regular["excess"]) == (True, 1)
    assert regular["calculation"] == INTAKE_EXCEEDED_TEXT
    assert [c["category"] for c in over.result["categories"] if c["exceeded"]] == [
        "REGULAR"
    ]
    population = seeded_session.scalar(
        select(IntakeAssessment).where(
            IntakeAssessment.batch_id
            == scenario.batches["06_cedar_missing_population.csv"]
        )
    )
    assert population is not None
    assert population.result["missing_population"]["calculation"] == SHORTFALL_TEXT

    # Integrity signals: all expected, none unexplained.
    signals = summary["integrity_signals"]
    assert (signals["false_positives"], signals["false_negatives"]) == (0, 0)
    assert signals["true_positives"] == 11

    # Report and analytics reconcile exactly.
    reconciliation = summary["reconciliation"]
    assert reconciliation["total"] >= 15
    assert reconciliation["equal"] == reconciliation["total"]
    assert all(r["equal"] for r in reconciliation["details"])

    # Authorization spot checks, repeated here independently of the runner.
    batch = scenario.batches["02_cedar_apaar_issues.csv"]
    cedar_case = scenario.ctx["case_id"]
    for alias, method, path, status in (
        ("demo_manager", "GET", "/api/v1/cases", 403),
        ("maple_admin", "GET", f"/api/v1/submissions/{batch}", 404),
        ("maple_admin", "GET", f"/api/v1/cases/{cedar_case}", 404),
        ("willow_admin", "GET", f"/api/v1/submissions/{batch}/rows", 404),
        ("demo_verifier", "GET", "/api/v1/cases", 403),
        ("demo_verifier", "GET", "/api/v1/audit/events", 403),
        ("demo_verifier", "GET", "/api/v1/records", 403),
        ("cedar_admin", "GET", "/api/v1/audit/events", 403),
    ):
        response = signed_client.request(method, path, headers=scenario.headers(alias))
        assert response.status_code == status, (alias, path, response.status_code)
    assert summary["authorization"]["passed"] == summary["authorization"]["total"]
    assert summary["authorization"]["total"] >= 12

    # Raw APAAR values never reach any response or any stored audit event.
    raw = raw_apaar_values(MANIFEST)
    assert len(raw) == 20 and len(scenario.transcript) > 100
    assert scenario.leaks() == []
    assert summary["raw_apaar_in_responses"]["leaks"] == []
    for label, _, text in scenario.transcript:
        assert not any(value in text for value in raw), label
    audit = seeded_session.scalars(select(AuditEvent.metadata_redacted)).all()
    assert audit
    audit_text = json.dumps(audit)
    assert not any(value in audit_text for value in raw)
    export = signed_client.get(
        "/api/v1/audit/events?page_size=100", headers=scenario.headers("demo_officer")
    )
    assert export.status_code == 200
    assert not any(value in export.text for value in raw)


def test_scenario_is_rerunnable_without_false_positives(
    signed_client: TestClient,
) -> None:
    first = Scenario(signed_client)
    assert first.run() is True
    second = Scenario(signed_client)
    # Same app instance, so reuse the demo sessions (the login limiter is per app).
    second.tokens = dict(first.tokens)
    assert second.run() is True
    summary: dict[str, Any] = second.summary()
    assert summary["fresh_database"] is False
    assert summary["workflow"]["steps_passed"] == 12
    accuracy = summary["validation_accuracy"]
    assert (accuracy["false_positives"], accuracy["false_negatives"]) == (0, 0)
    # Earlier-run history is reported separately, never hidden or mis-scored.
    assert accuracy["carryover_info_findings"] >= 1
    assert summary["raw_apaar_in_responses"]["leaks"] == []


def test_failed_expectation_is_reported_not_hidden(
    signed_client: TestClient,
) -> None:
    broken = json.loads(json.dumps(MANIFEST))
    broken["files"]["01_cedar_valid.csv"]["findings"] = [
        {
            "rule_code": "INVALID_APAAR_FORMAT",
            "discriminator": "apaar_id",
            "source_rows": [2],
            "severity": "WARNING",
            "classification": "REVIEWABLE",
        }
    ]
    scenario = Scenario(signed_client, broken)
    assert scenario.run() is False
    summary = scenario.summary()
    assert summary["passed"] is False
    assert summary["validation_accuracy"]["false_negatives"] == 1
    assert any(s["status"] == "FAIL" for s in summary["steps"])


def test_leak_scanner_detects_a_raw_value() -> None:
    # Positive control: the "0 leaks" result is only meaningful if the scan can fail.
    scenario = Scenario(httpx2.Client(base_url="http://localhost"))
    assert scenario.leaks() == []
    scenario.transcript.append(("GET /x", 200, "student 900000100001 listed"))
    assert scenario.leaks() == ["GET /x -> 200"]
