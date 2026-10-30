"""Report exports: authorization, reconciliation with analytics, formats, audit."""

import csv
import io
import re
import zlib
from collections import defaultdict
from collections.abc import Callable
from datetime import timedelta
from typing import Any

import pytest
from conftest import pin_clock
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_analytics import (
    CEDAR,
    MAPLE,
    NOW,
    YEAR,
    activate,
    add_case,
    get,
)
from test_validation import csv_data, run_validation, upload
from test_verification import one_record, verify

from app.modules.audit.models import AuditEvent
from app.modules.institutions.models import Institution
from app.modules.reports.renderers import latin1
from app.modules.reports.schemas import PROVENANCE

Headers = Callable[[str], dict[str, str]]
URL = "/api/v1/reports"
AGGREGATE = (
    "enrolment-summary",
    "intake-utilisation",
    "submission-completeness",
    "resolution-times",
    "validation-findings",
)
Tables = dict[str, list[dict[str, str]]]


def export(
    client: TestClient, kind: str, headers: dict[str, str], fmt: str = "csv", **q: Any
) -> Any:
    return client.get(f"{URL}/{kind}", headers=headers, params={"format": fmt, **q})


def parse_csv(text: str) -> tuple[dict[str, str], Tables]:
    """Metadata ('# key' rows) and each titled table as a list of dict rows."""
    metadata: dict[str, str] = {}
    tables: Tables = {}
    rows = list(csv.reader(io.StringIO(text)))
    index = 0
    while index < len(rows):
        row = rows[index]
        index += 1
        if not row:
            continue
        if row[0].startswith("# ") and len(row) == 2:
            metadata[row[0][2:]] = row[1]
        elif row[0].startswith("# "):
            title, header = row[0][2:], rows[index]
            index += 1
            body: list[dict[str, str]] = []
            while index < len(rows) and rows[index]:
                body.append(dict(zip(header, rows[index], strict=True)))
                index += 1
            tables[title] = body
    return metadata, tables


def csv_report(client: TestClient, kind: str, headers: dict[str, str], **q: Any) -> Any:
    response = export(client, kind, headers, **q)
    assert response.status_code == 200, response.text
    return parse_csv(response.text)


def table_with(tables: Tables, prefix: str) -> list[dict[str, str]]:
    return next(rows for title, rows in tables.items() if title.startswith(prefix))


def number(value: str) -> int | None:
    return int(value) if value != "" else None


def pdf_text(content: bytes) -> str:
    """Text of all (flate-compressed) page content streams."""
    out = []
    # Capture up to "endstream" and let the decompressor ignore the trailing line
    # break: a stream whose last byte is 0x0d would otherwise lose that byte to the
    # separator and fail to decompress (the timestamp makes the bytes vary per run).
    for raw in re.findall(rb"stream\r?\n(.*?)endstream", content, re.DOTALL):
        try:
            out.append(zlib.decompressobj().decompress(raw).decode("latin-1"))
        except zlib.error:
            continue
    return "\n".join(out)


# ------------------------------------------------------------------ authorization


@pytest.mark.parametrize(
    ("kind", "expected"),
    [
        *((kind, {"officer": 200, "manager": 200, "admin": 200}) for kind in AGGREGATE),
        ("open-cases", {"officer": 200, "manager": 403, "admin": 200}),
        ("verification-activity", {"officer": 200, "manager": 403, "admin": 403}),
    ],
)
def test_role_matrix(
    client: TestClient, auth_headers: Headers, kind: str, expected: dict[str, int]
) -> None:
    assert export(client, kind, {}).status_code == 401
    verifier = export(client, kind, auth_headers("demo_verifier"))
    assert verifier.status_code == 403
    assert verifier.json()["error"]["code"] == "FORBIDDEN_ROLE"
    for role, alias in (
        ("officer", "demo_officer"),
        ("manager", "demo_manager"),
        ("admin", "cedar_admin"),
    ):
        response = export(client, kind, auth_headers(alias))
        assert response.status_code == expected[role], (kind, role)


def test_invalid_kind_format_and_filters_are_rejected(
    client: TestClient, auth_headers: Headers
) -> None:
    officer = auth_headers("demo_officer")
    assert export(client, "unknown", officer).status_code == 422
    assert export(client, "open-cases", officer, fmt="docx").status_code == 422
    assert (
        export(client, "enrolment-summary", officer, academic_year="x").status_code
        == 422
    )


def test_institution_admin_is_limited_to_own_institution(
    client: TestClient, auth_headers: Headers
) -> None:
    cedar = auth_headers("cedar_admin")
    activate(client, cedar, [("REGULAR", "ACTIVE")])
    activate(client, auth_headers("maple_admin"), [("REGULAR", "ACTIVE")], "maple")
    other = export(client, "enrolment-summary", cedar, institution_id=str(MAPLE))
    assert other.status_code == 404
    _, tables = csv_report(client, "enrolment-summary", cedar)
    assert {r["institution"] for r in table_with(tables, "Programmes")} == {
        "Demo Cedar Institute"
    }
    meta, _ = csv_report(client, "enrolment-summary", cedar)
    assert meta["Filters applied"] == f"institution_id={CEDAR}"


# ------------------------------------------------------------------ layout


def test_csv_layout_headers_and_provenance(
    client: TestClient, application: FastAPI, auth_headers: Headers
) -> None:
    pin_clock(application, NOW)
    response = export(
        client,
        "enrolment-summary",
        auth_headers("demo_officer"),
        academic_year=YEAR,
        region="DEMO_REGION_1",
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["content-disposition"] == (
        'attachment; filename="enrolment-summary-20261001T120000Z.csv"'
    )
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    lines = response.text.splitlines()
    assert lines[0].startswith("# Report,") and "# Provenance" in response.text
    meta, tables = parse_csv(response.text)
    assert meta["Generated at (UTC)"] == NOW.isoformat()
    assert meta["Generated by"] == "AICTE_OFFICER (demo_officer)"
    assert "@" not in response.text
    assert meta["Filters applied"] == "academic_year=2026-27; region=DEMO_REGION_1"
    assert meta["Academic year context"] == "Academic year 2026-27"
    assert "current activated snapshot" in meta["Coverage"]
    assert meta["Provenance"] == PROVENANCE
    assert "SIMULATED / DEMO DATA" in meta["Provenance"]
    assert list(tables) == ["Totals by academic year", "Programmes"]


def test_missing_values_are_empty_never_zero(
    client: TestClient, auth_headers: Headers
) -> None:
    officer = auth_headers("demo_officer")
    _, tables = csv_report(client, "enrolment-summary", officer, academic_year=YEAR)
    for row in table_with(tables, "Totals")[:1] + table_with(tables, "Programmes"):
        assert row["status"] == "NO_CURRENT_SNAPSHOT"
        assert row["declared_records"] == "" and row["active_records"] == ""
    _, tables = csv_report(client, "intake-utilisation", officer, academic_year=YEAR)
    assert {r["current_active"] for r in table_with(tables, "Programme")} == {""}
    totals = table_with(tables, "Totals")
    assert {r["current_active"] for r in totals} == {""}
    assert {r["rows_with_known_counts"] for r in totals} == {"0"}

    response = export(client, "enrolment-summary", officer, "pdf", academic_year=YEAR)
    assert "No data" in pdf_text(response.content)
    workbook = load_workbook(
        io.BytesIO(export(client, "enrolment-summary", officer, "xlsx").content)
    )
    data = workbook["Data"]
    header = [c.value for c in data[2]]
    assert data.cell(row=3, column=header.index("active_records") + 1).value is None


# ------------------------------------------------------------------ reconciliation


def test_enrolment_summary_reconciles_with_analytics(
    client: TestClient, auth_headers: Headers
) -> None:
    cedar, maple = auth_headers("cedar_admin"), auth_headers("maple_admin")
    activate(
        client,
        cedar,
        [("REGULAR", "ACTIVE")] * 3 + [("NRI", "ACTIVE"), ("OCI", "WITHDRAWN")],
    )
    activate(client, maple, [("REGULAR", "ACTIVE")] * 2 + [("OCI", "ACTIVE")], "maple")
    officer = auth_headers("demo_officer")
    for filters in ({"academic_year": YEAR}, {}, {"region": "DEMO_REGION_1"}):
        meta, tables = csv_report(client, "enrolment-summary", officer, **filters)
        totals = {r["academic_year"]: r for r in table_with(tables, "Totals")}
        analytics = get(client, "enrolment-trends", officer, **filters)
        for row in analytics["data"]:
            mine = totals[row["academic_year"]]
            assert number(mine["declared_records"]) == row["declared_records"]
            assert number(mine["active_records"]) == row["active_records"]
            assert mine["status"] == row["status"]
            for category, counts in (row["categories"] or {}).items():
                assert number(mine[f"active_{category}"]) == counts["active_records"]
        assert len(totals) == len(analytics["data"])
        programmes = table_with(tables, "Programmes")
        known = [r for r in programmes if r["active_records"] != ""]
        year_total = totals.get(YEAR)
        if year_total and filters.get("region") is None:
            assert sum(int(r["active_records"]) for r in known) == int(
                year_total["active_records"]
            )
        coverage = analytics["coverage"]
        assert f"{coverage['programmes_total']} programmes" in meta["Coverage"]


def test_intake_utilisation_reconciles_with_analytics(
    client: TestClient, auth_headers: Headers
) -> None:
    cedar, maple = auth_headers("cedar_admin"), auth_headers("maple_admin")
    activate(client, cedar, [("REGULAR", "ACTIVE")] * 3 + [("NRI", "ACTIVE")])
    activate(client, maple, [("REGULAR", "ACTIVE")] * 2, "maple")
    officer = auth_headers("demo_officer")
    analytics = get(client, "intake-utilisation", officer, academic_year=YEAR)
    meta, tables = csv_report(client, "intake-utilisation", officer, academic_year=YEAR)
    rows = table_with(tables, "Programme-year")
    assert len(rows) == len(analytics["data"])
    for mine, theirs in zip(rows, analytics["data"], strict=True):
        assert mine["category"] == theirs["category"]
        assert number(mine["current_active"]) == theirs["current_active"]
        assert number(mine["approved_seats"]) == theirs["approved_seats"]
    expected: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for row in analytics["data"]:
        if row["status"] == "OK":
            expected[row["category"]][0] += row["approved_seats"] or 0
            expected[row["category"]][1] += row["current_active"]
    for total in table_with(tables, "Totals"):
        if total["category"] in expected:
            approved, active = expected[total["category"]]
            assert number(total["approved_seats"]) == approved
            assert number(total["current_active"]) == active
    assert "never pooled" in meta["Coverage"]


def test_submission_completeness_reconciles_with_analytics(
    client: TestClient, auth_headers: Headers
) -> None:
    admin, officer = auth_headers("cedar_admin"), auth_headers("demo_officer")
    activate(client, admin, [("REGULAR", "ACTIVE")] * 3)
    upload(client, admin, csv_data())  # a pending candidate
    analytics = get(client, "submission-completeness", officer, academic_year=YEAR)
    _, tables = csv_report(
        client, "submission-completeness", officer, academic_year=YEAR
    )
    totals = {r["region"]: r for r in table_with(tables, "Totals")}
    for rollup in analytics["rollups_by_region"]:
        mine = totals[rollup["region"]]
        for name, value in rollup.items():
            if name != "region":
                assert int(mine[name]) == value, name
    all_regions = totals["ALL REGIONS"]
    for name in ("institutions_total", "with_current_snapshot", "pending_candidates"):
        assert int(all_regions[name]) == sum(
            r[name] for r in analytics["rollups_by_region"]
        )
    rows = table_with(tables, "Programme-year")
    assert len(rows) == analytics["pagination"]["total"]
    cedar_row = next(r for r in rows if r["institution"] == "Demo Cedar Institute")
    assert cedar_row["pending_candidate_count"] == "1"


def test_case_reports_reconcile_with_case_performance(
    client: TestClient,
    application: FastAPI,
    auth_headers: Headers,
    seeded_session: Session,
) -> None:
    pin_clock(application, NOW)
    opened = NOW - timedelta(days=20)
    add_case(seeded_session, CEDAR, opened_at=opened, due_at=NOW + timedelta(days=1))
    add_case(seeded_session, CEDAR, opened_at=opened, due_at=NOW - timedelta(days=1))
    add_case(
        seeded_session,
        CEDAR,
        opened_at=opened,
        due_at=NOW - timedelta(days=7),
        status="UNDER_REVIEW",
        priority="HIGH",
        severity="HIGH",
    )
    add_case(
        seeded_session,
        MAPLE,
        opened_at=opened,
        due_at=NOW - timedelta(days=30),
        status="RESOLVED",
        resolved_at=NOW - timedelta(days=5),
    )
    officer = auth_headers("demo_officer")
    for filters in ({}, {"institution_id": str(CEDAR)}, {"priority": "HIGH"}):
        performance = get(client, "case-performance", officer, **filters)
        _, tables = csv_report(client, "open-cases", officer, **filters)
        totals = table_with(tables, "Open case totals")[0]
        expected = performance["totals"]
        assert int(totals["open"]) == expected["open"]
        assert int(totals["overdue"]) == expected["overdue"]
        assert int(totals["escalated"]) == expected["escalated"]
        rows = table_with(tables, "Open cases")
        assert len(rows) == expected["open"]
        assert sum(r["overdue"] == "true" for r in rows) == expected["overdue"]
        assert sum(r["escalated"] == "true" for r in rows) == expected["escalated"]
        flags = [r["overdue"] for r in rows]
        assert flags == sorted(flags, reverse=True)  # overdue first
        _, tables = csv_report(client, "resolution-times", officer, **filters)
        totals = table_with(tables, "Case totals")[0]
        assert int(totals["cases_total"]) == expected["all"]
        assert int(totals["open"]) == expected["open"]
        stats = {r["priority"]: r for r in table_with(tables, "Resolution")}
        for item in performance["resolution_days"]:
            mine = stats[item["priority"]]
            assert int(mine["resolved_count"]) == item["resolved_count"]
            if item["mean_days"] is None:
                assert mine["mean_days"] == ""
            else:
                assert float(mine["mean_days"]) == item["mean_days"]
    rows = table_with(csv_report(client, "open-cases", officer)[1], "Open cases")
    assert all(len(r["case_ref"]) == 8 for r in rows)
    assert set(rows[0]) >= {"institution", "status", "priority", "severity", "due_at"}
    # No resolved case yet for HIGH: statistics are empty, not zero.
    _, tables = csv_report(client, "resolution-times", officer, priority="HIGH")
    high = {r["priority"]: r for r in table_with(tables, "Resolution")}["HIGH"]
    assert high["resolved_count"] == "0" and high["mean_days"] == ""


def test_institution_admin_open_cases_only_own(
    client: TestClient,
    application: FastAPI,
    auth_headers: Headers,
    seeded_session: Session,
) -> None:
    pin_clock(application, NOW)
    add_case(seeded_session, CEDAR, opened_at=NOW - timedelta(days=3))
    add_case(seeded_session, MAPLE, opened_at=NOW - timedelta(days=3))
    _, tables = csv_report(client, "open-cases", auth_headers("maple_admin"))
    assert {r["institution"] for r in table_with(tables, "Open cases")} == {
        "Demo Maple Institute"
    }
    assert (
        export(
            client, "open-cases", auth_headers("maple_admin"), institution_id=str(CEDAR)
        ).status_code
        == 404
    )


def test_validation_findings_totals_match_the_findings_api(
    client: TestClient, auth_headers: Headers
) -> None:
    admin, officer = auth_headers("cedar_admin"), auth_headers("demo_officer")
    batch = upload(client, admin, csv_data(apaar_id="12"))
    meta, tables = csv_report(client, "validation-findings", officer)
    assert "0 of 1 latest completed submissions" in meta["Coverage"]
    assert {r["findings"] for r in table_with(tables, "Totals")} == {""}  # unknown
    run_validation(client, batch, officer)
    api = client.get(
        f"/api/v1/submissions/{batch}/findings",
        headers=officer,
        params={"page_size": 100},
    ).json()
    _, tables = csv_report(client, "validation-findings", officer, academic_year=YEAR)
    totals = {r["severity"]: int(r["findings"]) for r in table_with(tables, "Totals")}
    assert totals["ALL"] == api["pagination"]["total"] > 0
    assert totals["ALL"] == sum(
        int(r["findings"]) for r in table_with(tables, "Findings")
    )
    by_rule: dict[str, int] = defaultdict(int)
    for item in api["data"]:
        by_rule[item["rule_code"]] += 1
    mine: dict[str, int] = defaultdict(int)
    for row in table_with(tables, "Findings"):
        mine[row["rule_code"]] += int(row["findings"])
    assert mine == by_rule
    assert "12" not in {c for r in table_with(tables, "Findings") for c in r.values()}


def test_verification_activity_counts_without_identifiers(
    client: TestClient, auth_headers: Headers, seeded_session: Session
) -> None:
    record = one_record(client, auth_headers, seeded_session, "900000000001")
    officer = auth_headers("demo_officer")
    assert (
        verify(client, record, officer, purpose="CASE 123456789012 review").status_code
        == 201
    )
    assert verify(client, record, officer).status_code == 201
    logged = client.post(
        "/api/v1/credentials/verify",
        headers=auth_headers("demo_verifier"),
        json={"token": "synthetic", "purpose": "ADMISSION_CHECK", "consent_ref": "c1"},
    )
    assert logged.status_code == 200
    response = export(client, "verification-activity", officer)
    meta, tables = parse_csv(response.text)
    totals = {
        (r["source"], r["status"]): int(r["attempts"])
        for r in table_with(tables, "Totals")
    }
    assert totals[("SIMULATED_APAAR_VERIFICATION", "SIMULATED_MATCH")] == 2
    assert totals[("ALL", "ALL")] == sum(
        n for (source, _), n in totals.items() if source != "ALL"
    )
    attempts = table_with(tables, "Attempts")
    assert sum(int(r["attempts"]) for r in attempts) == totals[("ALL", "ALL")]
    assert any("[MASKED]" in r["purpose"] for r in attempts)
    assert not re.search(r"\d{12}", response.text)
    assert "DEMO-ONE" not in response.text and "Synthetic One" not in response.text
    assert "simulated provider" in meta["Coverage"]
    only_cedar = csv_report(
        client, "verification-activity", officer, institution_id=str(CEDAR)
    )[1]
    assert {r["source"] for r in table_with(only_cedar, "Attempts")} == {
        "SIMULATED_APAAR_VERIFICATION"
    }


# ------------------------------------------------------------------ formats


def test_xlsx_has_metadata_and_data_sheets_and_text_cells(
    client: TestClient, auth_headers: Headers, seeded_session: Session
) -> None:
    institution = seeded_session.get(Institution, CEDAR)
    assert institution is not None
    institution.name = '=HYPERLINK("https://example.invalid","click")'
    seeded_session.flush()
    officer = auth_headers("demo_officer")
    response = export(client, "enrolment-summary", officer, "xlsx", academic_year=YEAR)
    assert response.status_code == 200
    assert response.headers["content-type"].endswith("spreadsheetml.sheet")
    assert response.headers["content-disposition"].endswith('.xlsx"')
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Metadata", "Data"]
    keys = [row[0].value for row in workbook["Metadata"].iter_rows()]
    assert keys[0] == "Report" and "Provenance" in keys
    assert workbook["Metadata"]["B7"].value == PROVENANCE
    found = False
    for row in workbook["Data"].iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "HYPERLINK" in cell.value:
                found = True
                assert cell.data_type == "s" and cell.value.startswith("'=")
    assert found
    text = export(client, "enrolment-summary", officer, academic_year=YEAR).text
    assert "'=HYPERLINK" in text and ",=HYPERLINK" not in text


def test_pdf_is_paginated_with_repeating_header_and_footer(
    client: TestClient,
    application: FastAPI,
    auth_headers: Headers,
    seeded_session: Session,
) -> None:
    pin_clock(application, NOW)
    seeded_session.get(Institution, CEDAR).name = "Zürich ≠ 漢 demo"  # type: ignore[union-attr]
    for index in range(150):
        add_case(
            seeded_session,
            CEDAR,
            opened_at=NOW - timedelta(days=index % 40),
            due_at=NOW - timedelta(days=index % 5),
        )
    response = export(client, "open-cases", auth_headers("demo_officer"), "pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    pages = len(re.findall(rb"/Type /Page(?!s)", response.content))
    assert pages >= 3
    text = pdf_text(response.content)
    assert text.count("case ref") >= pages  # header repeats on every page
    # fpdf2 emits the page-count alias as a separate text run.
    assert re.search(rf"\(Page {pages} of \) Tj \({pages}\)", text)
    assert "(Page 1 of )" in text
    assert text.count("SIMULATED / DEMO DATA") >= pages  # footer provenance
    assert "Z\\374rich" in text or "Zürich" in text  # Latin-1 survives
    assert latin1("a—b 漢") == "a-b ?"


# ------------------------------------------------------------------ limits & audit


def test_report_too_large_is_refused_and_not_audited(
    client: TestClient,
    application: FastAPI,
    auth_headers: Headers,
    seeded_session: Session,
) -> None:
    for _ in range(3):
        add_case(seeded_session, CEDAR, opened_at=NOW - timedelta(days=1))
    application.state.settings = application.state.settings.model_copy(
        update={"report_max_rows": 2}
    )
    response = export(client, "open-cases", auth_headers("demo_officer"))
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "REPORT_TOO_LARGE"
    assert "Content-Disposition" not in response.headers
    assert (
        seeded_session.scalar(
            select(AuditEvent).where(AuditEvent.action == "REPORT_EXPORTED")
        )
        is None
    )


def test_export_is_audited_without_data(
    client: TestClient, auth_headers: Headers, seeded_session: Session
) -> None:
    assert (
        export(
            client,
            "enrolment-summary",
            auth_headers("demo_officer"),
            "xlsx",
            academic_year=YEAR,
            institution_id=str(CEDAR),
        ).status_code
        == 200
    )
    event = seeded_session.scalar(
        select(AuditEvent).where(AuditEvent.action == "REPORT_EXPORTED")
    )
    assert event is not None and event.entity_type == "REPORT"
    assert event.actor_id is not None and event.institution_id == CEDAR
    details = event.metadata_redacted
    assert details["kind"] == "enrolment-summary" and details["format"] == "xlsx"
    assert details["filter_keys"] == "academic_year,institution_id"
    assert details["institution_id"] == str(CEDAR) and details["row_count"] == "2"
    assert "Cedar" not in str(details) and "DEMO-COMP" not in str(details)
    admin_event = client.get(
        "/api/v1/audit/events",
        headers=auth_headers("demo_officer"),
        params={"action": "REPORT_EXPORTED"},
    )
    assert admin_event.json()["pagination"]["total"] == 1


def test_default_filters_and_case_filters_are_echoed(
    client: TestClient, auth_headers: Headers
) -> None:
    officer = auth_headers("demo_officer")
    meta, _ = csv_report(client, "open-cases", officer, priority="HIGH", region="R")
    assert meta["Filters applied"] == "region=R; priority=HIGH"
    meta, _ = csv_report(client, "resolution-times", officer)
    assert meta["Filters applied"] == "none"
